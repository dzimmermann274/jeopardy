#!/usr/bin/env python3
"""Build test-data/filled-test.xlsx — a filled workbook the way Google's xlsx
export delivers it (formulas already computed to literal values). Exercises:
bank tabs skipped/parsed, UNKNOWN answers, image IDs (Drive + direct URLs),
an "ID not found" error cell, blank questions (missing tiles), duplicate
values within a category, teams/players/title, and Final Jeopardy."""
import openpyxl
from pathlib import Path

wb = openpyxl.Workbook()

ws = wb.active
ws.title = "📖 READ ME FIRST"
ws["B2"] = "instructions here (ignored by the game)"

s = wb.create_sheet("⚙️ Game Setup")
s["B4"] = "Game title (shown on the big screen)"; s["C4"] = "Zimmermann Family Feud-pardy"
s["B7"] = "Team name"; s["C7"] = "Players — optional, separated by commas (shown on the TV)"
s["B8"] = "Team Grandma"; s["C8"] = "Teena, Danny, Amy"
s["B9"] = "Team Bach";    s["C9"] = "Paul, Sarah"
s["B16"] = "FINAL JEOPARDY — optional."
s["B17"] = "Final Jeopardy category"; s["C17"] = "Family History"
s["B18"] = "Final Jeopardy question"; s["C18"] = "This is the final question."
s["B19"] = "Final Jeopardy answer (or UNKNOWN)"; s["C19"] = "The final answer"

qb = wb.create_sheet("🗂 Question Bank")
# banner row mimicking the real template (must NOT be mistaken for the header)
qb["A1"] = "🗂 QUESTION BANK — write ALL your questions here... Each row's ID number is how you drop the question into a category tab."
qb["A3"] = "ID"; qb["B3"] = "QUESTION"; qb["C3"] = "ANSWER — or type UNKNOWN"; qb["D3"] = "Image ID (optional)"
bank = [
    (1, "Bank question one?", "Bank answer one", ""),
    (2, "Bank question two?", "UNKNOWN", ""),
    (3, "Bank question three with a picture?", "Bank answer three", 1),
]
for i, (qid, q, a, img) in enumerate(bank):
    r = 4 + i
    qb[f"A{r}"] = qid; qb[f"B{r}"] = q; qb[f"C{r}"] = a; qb[f"D{r}"] = img

ib = wb.create_sheet("🖼 Image Bank")
# banner + how-to rows mimicking the real template — a single cell containing
# both "ID" and "links" once fooled the header detection (regression guard)
ib["A1"] = "🖼 IMAGE BANK — paste picture links here. Attach one to a question by putting its Image ID next to the question."
ib["A2"] = "Easiest way: upload to Google Drive → Share → Anyone with the link → Copy link → paste below."
ib["A4"] = "Image ID"; ib["B4"] = "Image link (paste the URL here)"; ib["C4"] = "What it shows"
ib["A5"] = 1; ib["B5"] = "https://drive.google.com/file/d/1AbCdEfGhIjKlMnOpQrStUvWxYz123456/view?usp=sharing"; ib["C5"] = "grandma photo"
ib["A6"] = 2; ib["B6"] = "https://upload.wikimedia.org/wikipedia/commons/1/15/Cat_August_2010-4.jpg"; ib["C6"] = "a cat"

def cat_tab(name, rows, dd=None):
    ws = wb.create_sheet(name)
    ws["A3"] = "Value"; ws["B3"] = "Question ID (from the 🗂 bank)"
    ws["C3"] = "QUESTION — auto-fills from the ID"; ws["D3"] = "ANSWER — auto-fills. UNKNOWN = host types it live"
    ws["E3"] = "Image ID (auto)"
    for i, (val, q, a, img) in enumerate(rows):
        r = 4 + i
        ws[f"A{r}"] = f"${val}" if val else ""
        ws[f"C{r}"] = q; ws[f"D{r}"] = a; ws[f"E{r}"] = img
    if dd:
        ws["A10"] = "💰 Daily Double (optional): type ONE dollar amount →"
        ws["D10"] = dd
    return ws

# Category A: full 5, one UNKNOWN answer, one image via bank id, DD at 800
cat_tab("Grandma Lore", [
    (200, "Bank question one?", "Bank answer one", ""),
    (400, "Bank question two?", "UNKNOWN", ""),
    (600, "Bank question three with a picture?", "Bank answer three", 1),
    (800, "Direct-typed question?", "Direct answer", ""),
    (1000, "Another question?", "Another answer", 2),
], dd=800)

# Category B: gaps (blank question at 400/1000) + an unresolved ID cell
cat_tab("Holes & Errors", [
    (200, "B two hundred?", "B two hundred answer", ""),
    (400, "", "", ""),
    (600, "⚠ ID not found", "⚠ ID not found", ""),
    (800, "B eight hundred?", "B eight hundred answer", ""),
    (1000, "", "", ""),
])

# Category C: duplicate values (two $600 rows — writer typo)
cat_tab("Duplicate Values", [
    (200, "C first?", "C first answer", ""),
    (600, "C second?", "C second answer", ""),
    (600, "C third (duplicate value)?", "C third answer", ""),
    (1000, "C fourth?", "C fourth answer", ""),
])

# Category D: review-regression cases — a clue that MENTIONS "Daily Double",
# a blank answer (-> live-typed), an "Unknown?" variant, a bad Image ID
cat_tab("Tricky Cases", [
    (200, "This game show is famous for its Daily Double wagers.", "Jeopardy!", ""),
    (400, "Question with no answer typed?", "", ""),
    (600, "Question with unknown-variant answer?", "Unknown?", ""),
    (800, "Question with a bad image id?", "Some answer", 99),
], dd=600)

# untouched category tab -> must not appear
empty = wb.create_sheet("Category 6")
empty["A3"] = "Value"; empty["B3"] = "Question ID (from the 🗂 bank)"
empty["C3"] = "QUESTION — auto-fills from the ID"; empty["D3"] = "ANSWER — auto-fills"

out = Path(__file__).resolve().parent.parent / "test-data" / "filled-test.xlsx"
out.parent.mkdir(parents=True, exist_ok=True)
wb.save(out)
print(f"wrote {out}")
