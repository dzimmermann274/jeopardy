#!/usr/bin/env python3
"""Generate the Jeopardy question workbook (template/Jeopardy Questions.xlsx).

Tabs: 📖 READ ME FIRST | ⚙️ Game Setup | 🗂 Question Bank | 🖼 Image Bank | Category 1-6

Layout contract (the game's js/data.js parser depends on it):
- Tabs named like "read me"/"instructions" and "game setup" are metadata;
  "question bank" and "image bank" are the banks; every OTHER tab is one
  category, named by its tab name.
- Question Bank: header row with ID | Question | Answer | Image ID; writers
  give every question a number so category tabs can reference it.
- Image Bank: header row with Image ID | link; maps image numbers to URLs.
- Category tabs: header row containing "Question" and "Answer" (the Question
  ID column header also contains the word Question, so the parser matches
  "question" NOT followed by "id"); data rows: Value | Question ID |
  Question | Answer | Image ID. The Question/Answer/Image cells hold VLOOKUP
  formulas keyed on the ID cell — Google computes them, and the game reads
  the computed values from the xlsx export. Typing directly over a formula
  also works. Blank question -> no tile.
- An answer of UNKNOWN (any case) means: no preset answer; the host types
  the answer live in the control panel during the game.
- A cell containing "Daily Double" marks a label; the first number found to
  its right names the dollar level that becomes the Daily Double.
- The Game Setup tab supplies: "Game title" label -> title cell; a teams
  table headed "Team name" / "Players"; and optional "Final Jeopardy
  category/question/answer" label -> value rows.

Rerun any time:  python3 tools/make_template.py
"""
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from pathlib import Path

NAVY = "1F2A6B"          # header bars
BOARD_BLUE = "2437D9"    # accents / banner
GOLD = "F5C518"          # highlights
LABEL_GRAY = "ECEEF4"    # label cells
AUTO_BLUE = "EAF0FB"     # auto-fill (formula) cells
ENTRY_BORDER = "4A7DFF"  # entry cell border (blue = "type here")
TEXT_DARK = "20242F"

QB = "🗂 Question Bank"
IB = "🖼 Image Bank"

med_blue = Side(style="medium", color=ENTRY_BORDER)
thin_gray = Side(style="thin", color="B9C0D4")
ENTRY_BORDER_STYLE = Border(left=med_blue, right=med_blue, top=med_blue, bottom=med_blue)
AUTO_BORDER_STYLE = Border(left=thin_gray, right=thin_gray, top=thin_gray, bottom=thin_gray)
ENTRY_FILL = PatternFill("solid", fgColor="FFFFFF")
AUTO_FILL = PatternFill("solid", fgColor=AUTO_BLUE)
LABEL_FILL = PatternFill("solid", fgColor=LABEL_GRAY)
HEADER_FILL = PatternFill("solid", fgColor=NAVY)
GOLD_FILL = PatternFill("solid", fgColor=GOLD)
WHITE_FILL = PatternFill("solid", fgColor="FFFFFF")

F_TITLE = Font(name="Arial", size=22, bold=True, color=NAVY)
F_HEADER = Font(name="Arial", size=12, bold=True, color="FFFFFF")
F_SECTION = Font(name="Arial", size=14, bold=True, color=NAVY)
F_LABEL = Font(name="Arial", size=12, bold=True, color=TEXT_DARK)
F_BODY = Font(name="Arial", size=12, color=TEXT_DARK)
F_ENTRY = Font(name="Arial", size=13, color=TEXT_DARK)
F_AUTO = Font(name="Arial", size=13, italic=False, color="3A4358")
F_VALUE = Font(name="Arial", size=18, bold=True, color=NAVY)
F_ID = Font(name="Arial", size=14, bold=True, color=NAVY)

WRAP_MID = Alignment(wrap_text=True, vertical="center")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)


def entry(ws, cell, height=None, font=F_ENTRY, align=WRAP_MID):
    c = ws[cell]
    c.fill = ENTRY_FILL
    c.border = ENTRY_BORDER_STYLE
    c.font = font
    c.alignment = align
    if height:
        ws.row_dimensions[c.row].height = height
    return c


def auto_cell(ws, cell, formula):
    c = ws[cell]
    c.value = formula
    c.fill = AUTO_FILL
    c.border = AUTO_BORDER_STYLE
    c.font = F_AUTO
    c.alignment = WRAP_MID
    return c


def label(ws, cell, text, font=F_LABEL, fill=LABEL_FILL, align=WRAP_MID):
    c = ws[cell]
    c.value = text
    c.font = font
    c.fill = fill
    c.alignment = align
    return c


def make_instructions(wb):
    ws = wb.active
    ws.title = "📖 READ ME FIRST"
    ws.sheet_properties.tabColor = GOLD
    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = 108
    ws.sheet_view.showGridLines = False

    label(ws, "B2", "🎯 JEOPARDY QUESTION WORKBOOK", F_TITLE, WHITE_FILL)
    ws.row_dimensions[2].height = 34

    rows = [
        ("STEP 1 — WRITE YOUR QUESTIONS IN THE 🗂 QUESTION BANK", "section"),
        ("Put every question you come up with in the Question Bank tab, one per row — in any order. Each row has an ID number; that number is how you'll drop the question into a category later, so questions are easy to swap around.", "body"),
        ("Special answer: type UNKNOWN as the answer if there's no preset answer — the game host will type the real answer live during the game (great for crowd challenges and judgment calls).", "body"),
        ("Optional “Answer replaces question?” column: type YES to make that answer fill the whole screen when it's revealed (the question disappears — but a photo, if any, always stays). Leave it blank for a normal clue where the question and answer show together.", "body"),
        ("", "body"),
        ("STEP 2 — (OPTIONAL) ADD PICTURES IN THE 🖼 IMAGE BANK", "section"),
        ("Paste image links in the Image Bank tab, one per row, each with an Image ID. Easiest source: upload the picture to Google Drive → right-click it → Share → \"Anyone with the link\" → Copy link → paste here. Regular image links from the web work too.", "body"),
        ("Attach a picture to a question by putting its Image ID in the question's Image ID column (in the Question Bank, or on a category tab). The picture appears on the TV under the question. For TWO pictures on one question, put both Image IDs separated by a comma (e.g. 3, 7) — they show side by side.", "body"),
        ("", "body"),
        ("STEP 3 — BUILD THE CATEGORIES", "section"),
        ("Each colored \"Category\" tab is ONE category on the game board (up to 6). Double-click the tab and RENAME it — the tab's name is exactly what appears on the board. (\"Category 1\" → \"80s Movies\")", "body"),
        ("Next to each dollar amount, just type a Question ID from the bank — the question and answer fill in by themselves (blue-tinted cells). You can also ignore the bank and type a question straight into those cells.", "body"),
        ("Leave a row's ID blank → that tile won't appear on the board. Leave a whole tab untouched → that category won't appear at all. The game board mirrors this workbook exactly.", "body"),
        ("Optional: give a category a Daily Double by typing one dollar amount (like 800) in the gold box at the bottom of its tab.", "body"),
        ("", "body"),
        ("STEP 4 — GAME SETUP TAB", "section"),
        ("Fill in the game title, team names, and (optionally) each team's players — these show on the TV. There's also an optional Final Jeopardy section.", "body"),
        ("", "body"),
        ("WHEN YOU'RE FINISHED", "section"),
        ("Click the blue Share button (top right) → under \"General access\" choose \"Anyone with the link\" → Viewer → Copy link. Send that link to the game host — it's all they need.", "body"),
        ("🤫 DON'T show the host the answers (they play too!). The game keeps every answer hidden until it's revealed on the TV for everyone at once.", "body"),
        ("", "body"),
        ("WHERE DO I TYPE?", "section"),
        ("Type in the WHITE boxes with BLUE borders. Blue-tinted cells fill in automatically from the banks (you may type over them). Gray cells are labels — leave those alone.", "body"),
    ]
    r = 4
    for text, kind in rows:
        if kind == "section":
            label(ws, f"B{r}", text, F_SECTION, WHITE_FILL)
            ws.row_dimensions[r].height = 26
        else:
            label(ws, f"B{r}", text, F_BODY, WHITE_FILL)
            ws.row_dimensions[r].height = 34 if len(text) > 110 else 24
        r += 1


def make_setup(wb):
    ws = wb.create_sheet("⚙️ Game Setup")
    ws.sheet_properties.tabColor = "555B6E"
    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = 34
    ws.column_dimensions["C"].width = 70
    ws.sheet_view.showGridLines = False

    label(ws, "B2", "⚙️ GAME SETUP", F_TITLE, WHITE_FILL)
    ws.row_dimensions[2].height = 32

    label(ws, "B4", "Game title (shown on the big screen)")
    entry(ws, "C4", height=28)
    ws["C4"] = "Jeopardy!"

    label(ws, "B6", "TEAMS — up to 6. Blank rows are skipped.", F_SECTION, WHITE_FILL)
    ws.row_dimensions[6].height = 24
    label(ws, "B7", "Team name", F_HEADER, HEADER_FILL, CENTER)
    label(ws, "C7", "Players — optional, separated by commas (shown on the TV)", F_HEADER, HEADER_FILL, CENTER)
    ws.row_dimensions[7].height = 24
    for i in range(8, 14):
        entry(ws, f"B{i}", height=24)
        entry(ws, f"C{i}")

    label(ws, "B16", "FINAL JEOPARDY — optional. Leave blank for no final round.", F_SECTION, WHITE_FILL)
    ws.row_dimensions[16].height = 24
    label(ws, "B17", "Final Jeopardy category")
    entry(ws, "C17", height=26)
    label(ws, "B18", "Final Jeopardy question")
    entry(ws, "C18", height=60)
    label(ws, "B19", "Final Jeopardy answer (or UNKNOWN)")
    entry(ws, "C19", height=30)


def make_question_bank(wb):
    ws = wb.create_sheet(QB)
    ws.sheet_properties.tabColor = "1FA05A"
    ws.column_dimensions["A"].width = 8
    ws.column_dimensions["B"].width = 72
    ws.column_dimensions["C"].width = 42
    ws.column_dimensions["D"].width = 12
    ws.column_dimensions["E"].width = 34
    ws.sheet_view.showGridLines = False

    ws.merge_cells("A1:E1")
    label(ws, "A1",
          "🗂  QUESTION BANK — write ALL your questions here, one per row, in any order. "
          "Each row's ID number is how you drop the question into a category tab, so questions stay easy to swap.",
          Font(name="Arial", size=12, bold=True, color="FFFFFF"),
          PatternFill("solid", fgColor="1FA05A"), WRAP_MID)
    ws.row_dimensions[1].height = 40

    label(ws, "A3", "ID", F_HEADER, HEADER_FILL, CENTER)
    label(ws, "B3", "QUESTION", F_HEADER, HEADER_FILL, CENTER)
    label(ws, "C3", "ANSWER — or type UNKNOWN to have the host type it live", F_HEADER, HEADER_FILL, CENTER)
    label(ws, "D3", "Image ID (optional)", F_HEADER, HEADER_FILL, CENTER)
    label(ws, "E3", "Answer replaces question? Type YES to hide the question when the answer shows (photos always stay).", F_HEADER, HEADER_FILL, CENTER)
    ws.row_dimensions[3].height = 30

    for i in range(1, 61):
        r = 3 + i
        label(ws, f"A{r}", i, F_ID, LABEL_FILL, CENTER)
        entry(ws, f"B{r}", height=34)
        entry(ws, f"C{r}")
        entry(ws, f"D{r}", align=CENTER)
        entry(ws, f"E{r}", align=CENTER)


def make_image_bank(wb):
    ws = wb.create_sheet(IB)
    ws.sheet_properties.tabColor = "B0559F"
    ws.column_dimensions["A"].width = 10
    ws.column_dimensions["B"].width = 85
    ws.column_dimensions["C"].width = 34
    ws.sheet_view.showGridLines = False

    ws.merge_cells("A1:C1")
    label(ws, "A1",
          "🖼  IMAGE BANK — paste picture links here. Attach one to a question by putting its Image ID "
          "next to the question (in the 🗂 Question Bank or on a category tab).",
          Font(name="Arial", size=12, bold=True, color="FFFFFF"),
          PatternFill("solid", fgColor="B0559F"), WRAP_MID)
    ws.row_dimensions[1].height = 40

    ws.merge_cells("A2:C2")
    label(ws, "A2",
          "Easiest way: upload the picture to Google Drive → right-click → Share → \"Anyone with the link\" → Copy link → paste below. "
          "Regular image links from the web work too.",
          F_BODY, WHITE_FILL, WRAP_MID)
    ws.row_dimensions[2].height = 30

    label(ws, "A4", "Image ID", F_HEADER, HEADER_FILL, CENTER)
    label(ws, "B4", "Image link (paste the URL here)", F_HEADER, HEADER_FILL, CENTER)
    label(ws, "C4", "What it shows (just a note for you)", F_HEADER, HEADER_FILL, CENTER)
    ws.row_dimensions[4].height = 26

    for i in range(1, 31):
        r = 4 + i
        label(ws, f"A{r}", i, F_ID, LABEL_FILL, CENTER)
        entry(ws, f"B{r}", height=24)
        entry(ws, f"C{r}")


def make_category(wb, idx):
    ws = wb.create_sheet(f"Category {idx}")
    shades = ["2437D9", "3A4DE0", "1E2FBF", "4A5CE8", "16249B", "5B6CEF"]
    ws.sheet_properties.tabColor = shades[(idx - 1) % len(shades)]
    ws.column_dimensions["A"].width = 10
    ws.column_dimensions["B"].width = 14
    ws.column_dimensions["C"].width = 58
    ws.column_dimensions["D"].width = 38
    ws.column_dimensions["E"].width = 11
    ws.sheet_view.showGridLines = False

    ws.merge_cells("A1:E1")
    label(ws, "A1",
          f"✏️  This whole tab is ONE category. Double-click the tab name below (\"Category {idx}\") and rename it — "
          "the tab's name is the category title players see on the board. Type a Question ID from the 🗂 Question Bank "
          "next to each dollar amount and the rest fills in by itself (or just type a question directly).",
          Font(name="Arial", size=12, bold=True, color="FFFFFF"),
          PatternFill("solid", fgColor=BOARD_BLUE), WRAP_MID)
    ws.row_dimensions[1].height = 52

    label(ws, "A3", "Value", F_HEADER, HEADER_FILL, CENTER)
    label(ws, "B3", "Question ID (from the 🗂 bank)", F_HEADER, HEADER_FILL, CENTER)
    label(ws, "C3", "QUESTION — auto-fills from the ID (you can also type here)", F_HEADER, HEADER_FILL, CENTER)
    label(ws, "D3", "ANSWER — auto-fills. UNKNOWN = host types it live", F_HEADER, HEADER_FILL, CENTER)
    label(ws, "E3", "Image ID (auto)", F_HEADER, HEADER_FILL, CENTER)
    ws.row_dimensions[3].height = 40

    for i, value in enumerate([200, 400, 600, 800, 1000]):
        r = 4 + i
        label(ws, f"A{r}", f"${value}", F_VALUE, LABEL_FILL, CENTER)
        entry(ws, f"B{r}", height=64, font=F_ID, align=CENTER)
        auto_cell(ws, f"C{r}", f"=IF($B{r}=\"\",\"\",IFERROR(VLOOKUP($B{r},'{QB}'!$A:$D,2,FALSE),\"⚠ ID not found\"))")
        auto_cell(ws, f"D{r}", f"=IF($B{r}=\"\",\"\",IFERROR(VLOOKUP($B{r},'{QB}'!$A:$D,3,FALSE),\"⚠ ID not found\"))")
        auto_cell(ws, f"E{r}", f"=IF($B{r}=\"\",\"\",IFERROR(VLOOKUP($B{r},'{QB}'!$A:$D,4,FALSE),\"\"))")

    ws.merge_cells("A10:C10")
    label(ws, "A10",
          "💰 Daily Double (optional): type ONE dollar amount from above (e.g. 800) to make that question "
          "the Daily Double →",
          Font(name="Arial", size=12, bold=True, color=TEXT_DARK), GOLD_FILL, WRAP_MID)
    entry(ws, "D10", height=34)


def main():
    wb = openpyxl.Workbook()
    make_instructions(wb)
    make_setup(wb)
    make_question_bank(wb)
    make_image_bank(wb)
    for i in range(1, 7):
        make_category(wb, i)
    out = Path(__file__).resolve().parent.parent / "template" / "Jeopardy Questions.xlsx"
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    print(f"wrote {out}")


if __name__ == "__main__":
    main()
